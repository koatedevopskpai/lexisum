"""
In-memory DOCX / PDF export engine for LexiSum briefings.

Produces professional, formatted reports directly into memory buffers so no
temporary files are written to disk - ideal for both the Streamlit UI
(download buttons) and the CLI (saving to an output directory).
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, Optional


class DocumentExportEngine:
    """Generate DOCX and PDF briefings from a pipeline result."""

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------
    @staticmethod
    def create_docx(result: Dict[str, Any], doc_title: str = "Contract Executive Briefing") -> bytes:
        """Build a .docx report and return its bytes."""
        from docx import Document
        from docx.shared import Pt

        doc = Document()
        doc.add_heading(doc_title, level=0)

        doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC")
        _add_metadata(doc, result)

        # Extractive clauses
        extractive = result.get("extractive_summary", {})
        doc.add_heading("Verbatim Extractive Clauses", level=1)
        for i, sent in enumerate(extractive.get("sentences", []), 1):
            p = doc.add_paragraph()
            run = p.add_run(f"Clause {i}: ")
            run.bold = True
            p.add_run(sent)
        verified = extractive.get("compliance_verified")
        p = doc.add_paragraph()
        run = p.add_run(f"Verbatim audit: {'PASSED' if verified else 'FAILED'}")
        run.bold = True

        # Abstractive summary
        abstractive = result.get("abstractive_summary", {})
        doc.add_heading("Plain-Language Executive Summary", level=1)
        doc.add_paragraph(abstractive.get("text", ""))
        p = doc.add_paragraph()
        run = p.add_run(
            f"Structure audit: {'PASSED' if abstractive.get('structure_verified') else 'FAILED'}"
        )
        run.bold = True

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------
    @staticmethod
    def create_pdf(result: Dict[str, Any], doc_title: str = "Contract Executive Briefing") -> bytes:
        """Build a .pdf report and return its bytes."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "LexiTitle", parent=styles["Title"], textColor=colors.HexColor("#0f172a")
        )
        h1 = ParagraphStyle(
            "LexiH1", parent=styles["Heading1"], spaceBefore=12, textColor=colors.HexColor("#1d4ed8")
        )
        body = ParagraphStyle(
            "LexiBody",
            parent=styles["BodyText"],
            leading=14,
            spaceAfter=8,
        )

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1 * inch, rightMargin=1 * inch)
        story: list = []

        story.append(Paragraph(doc_title, title_style))
        story.append(Spacer(1, 8))
        story.append(
            Paragraph(
                f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC", styles["Italic"]
            )
        )
        story.append(Spacer(1, 12))

        extractive = result.get("extractive_summary", {})
        story.append(Paragraph("Verbatim Extractive Clauses", h1))
        for i, sent in enumerate(extractive.get("sentences", []), 1):
            story.append(Paragraph(f"<b>Clause {i}:</b> {_escape(sent)}", body))
        verified = extractive.get("compliance_verified")
        story.append(
            Paragraph(
                f"<b>Verbatim audit: {'PASSED' if verified else 'FAILED'}</b>", body
            )
        )
        story.append(Spacer(1, 8))

        abstractive = result.get("abstractive_summary", {})
        story.append(Paragraph("Plain-Language Executive Summary", h1))
        for para in abstractive.get("text", "").split("\n\n"):
            story.append(Paragraph(_escape(para), body))
        story.append(
            Paragraph(
                f"<b>Structure audit: {'PASSED' if abstractive.get('structure_verified') else 'FAILED'}</b>",
                body,
            )
        )

        doc.build(story)
        return buf.getvalue()


def _add_metadata(doc, result: Dict[str, Any]) -> None:
    """Add document metadata to a python-docx Document."""
    meta = result.get("metadata", {})
    if not meta:
        return
    words = meta.get("word_count")
    sents = meta.get("total_sentences")
    if words is not None or sents is not None:
        line = " ".join(
            str(x)
            for x in (
                f"~{words} words" if words is not None else None,
                f"{sents} sentences" if sents is not None else None,
            )
            if x
        )
        doc.add_paragraph(line)


def _escape(text: str) -> str:
    """Escape XML special characters for ReportLab paragraphs."""
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
